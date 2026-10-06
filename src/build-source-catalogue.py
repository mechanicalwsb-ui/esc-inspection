"""Read-only source extraction. Outputs only inside the inspection project."""
from pathlib import Path
import hashlib, json, re
from pypdf import PdfReader
from openpyxl import load_workbook

SOURCE = Path(r'C:\Users\autocad06\OneDrive - easternsugar.co.th\02.OneDriver วิศวกรรมจักรกล\02.ลูกหีบ-SP\03.เอกสาร\18.เอกสารบันทึก')
ROOT = Path(__file__).resolve().parent.parent

def field(label, kind='text', unit='', standard=''):
    return dict(label=label, kind=kind, unit=unit, standard=standard)
def numbers(labels, unit=''):
    return [field(x, 'number', unit) for x in labels.split('|')]

schemas = {
 'FM-ML01-ML-02': numbers('ความเร็วรอบเทอร์ไบน์','rpm') + numbers('FBV-X|FBV-Y|RBV-X|RBV-Y','µm') + [field('Axial Dis.', 'number', 'mm'),field('อัตราการไหลของไอดี','number','ton/hr')] + numbers('อุณหภูมิไอดี|อุณหภูมิไอเสีย|แบร์ริ่งเทอร์ไบน์หน้า|แบร์ริ่งเทอร์ไบน์หลัง|แบร์ริ่งเกียร์ขับหน้า|แบร์ริ่งเกียร์ขับหลัง|น้ำมันหล่อลื่นเข้า|น้ำมันหล่อลื่นออก|น้ำหล่อเย็นเข้า|น้ำหล่อเย็นออก|ตุ๊กตาเชร็ดเดอร์ฝั่งเทอร์ไบน์|ตุ๊กตาเชร็ดเดอร์ฝั่งตรงข้าม','°C') + numbers('แรงดันไอดี|แรงดันไอเสีย','kg/cm²') + numbers('แรงดันน้ำมันคอนโทรล|แรงดันน้ำมันแบร์ริ่ง','MPa') + numbers('แรงดันน้ำมันเกียร์|แรงดันน้ำมันตุ๊กตาเชร็ดเดอร์|Offline Filter','kg/cm²'),
 'FM-ML01-ML-03': [field('ลูกหีบชุดที่'),field('ยี่ห้อ/รุ่นจารบี')] + numbers('Off Time|Run Time','นาที') + numbers('แรงดันไฮดรอลิคฝั่ง A|แรงดันไฮดรอลิคฝั่ง B','kg/cm²') + [field('ตำแหน่งลูก/ฝั่ง A หรือ B'),field('เพลา','number','°C','<50'),field('ฝาบน','number','°C','25–40'),field('ฝาล่าง','number','°C','25–40'),field('น้ำเลี้ยง','number','°C','20–35')],
 'FM-ML01-ML-05': [field('เริ่มลดรอบ','time'),field('สิ้นสุดลดรอบ','time'),field('รวมเวลา','number','นาที'),field('สาเหตุ'),field('วิธีแก้ไข')],
 'FM-ML01-ML-06': [field('เริ่มหยุดหีบ','time'),field('เริ่มหีบ','time'),field('รวมเวลาหยุด','number','นาที'),field('สาเหตุ'),field('วิธีแก้ไข')],
 'FM-ML01-ML-07': [field('ลูกหีบชุดที่'),field('ยี่ห้อ/รุ่นจารบี'),field('รุ่นเครื่องเติม'),field('ถังที่'),field('ชนิดถัง (ใหม่/ร่วม)')] + numbers('Off Time|Run Time','นาที') + [field('เริ่มใช้','datetime-local'),field('หมดถัง','datetime-local')] + numbers('ก่อนใช้|หลังใช้|ใช้จารบี','kg'),
 'FM-ML01-ML-08': [field('ตำแหน่งลูกหีบที่ร้อน'),field('ตำแหน่งเพลาที่ร้อน/ฝั่ง')] ,
 'FM-ML01-ML-09': [field('จุดตรวจความสะอาด (มอเตอร์/กระบอก/หมายเลข)'),field('ความสะอาดและรอยรั่ว','check')] + [field('ระดับน้ำมันในถัง','number','%','>80')] + numbers('น้ำหล่อเย็นเข้า|น้ำหล่อเย็นออก|น้ำมันไฮดรอลิคเข้า|น้ำมันไฮดรอลิคออก|มอเตอร์ยกดั๊มพ์ 1|มอเตอร์ยกดั๊มพ์ 2|มอเตอร์ล็อคดั๊มพ์ 1|มอเตอร์ล็อคดั๊มพ์ 2','°C') + numbers('แรงดัน Main|แรงดัน Offline Filter','kg/cm²'),
 'FM-ML01-ML-10': [field('เครื่อง/จุดตรวจ/เบอร์สะพาน')] + [field(x,'number','°C','<80') for x in ['MT','G','DE','NDE']],
 'FM-ML01-ML-11': [field('สะพาน (เชรดเดอร์/High speed belt/ข้ามชุด 3–5)')] + [field(x,'number','°C','<80') for x in ['MT','G','DE','NDE']],
 'FM-ML01-ML-12': [field('ลำดับรายการ/ชื่อเครื่องจักรตามต้นฉบับ'),field('จุดตรวจ'),field('ค่าที่วัด','number'),field('หน่วย'),field('ค่าควบคุมตามต้นฉบับ'),field('ผลตรวจ','check')],
 'FM-ML01-ML-13': [field('รุ่นเครื่องชั่ง'),field('เริ่มเก็บ','time'),field('เก็บเสร็จ','time'),field('รวมเวลา','number','นาที')] + numbers('เชรดเดอร์ โซ่/สลิง/เป็คยันอ้อย|เชรดเดอร์ ใบมีด|เชรดเดอร์ อื่นๆ|High Speed Belt โซ่/สลิง/เป็คยันอ้อย|High Speed Belt ใบมีด|High Speed Belt อื่นๆ|รวมน้ำหนัก','kg'),
 'FM-ML01-ML-14': [field('จุดฆ่าเชื้อ (ลูกหีบ 1–4/Rotary)')] + numbers('ระยะเวลาเปิดน้ำร้อน|ระยะเวลาเปิดไอ','นาที') + [field('กากอ้อยสะสมใต้การ์ด','check'),field('เชื้อแบคทีเรีย','check'),field('ล้างลูกหีบ')],
 'FM-ML01-ML-15': [field('ลำดับขั้นตอนตามต้นฉบับ'),field('เลขเครื่องจักรที่รัน'),field('เวลาเริ่ม','time'),field('เวลาสิ้นสุด/ต่อเนื่อง'),field('ค่าที่บันทึก','number'),field('หน่วย'),field('ค่าควบคุม'),field('ผลตรวจ','check')],
 'FM-ML01-ML-16': [field('ดั๊มพ์เบอร์/ราง'),field('ทะเบียนรถ'),field('ชนิดสิ่งปนเปื้อน'),field('น้ำหนัก','number','kg'),field('เวลาตรวจพบ','time'),field('คนขับรถ'),field('รายละเอียดเอกสารแนบ')],
 'FM-ML01-ML-17': numbers('น้ำหนักต่อครั้ง|น้ำหนักพันยอด','kg'),
 'FM-ML01-ML-18': [field('ทะเบียนรถ'),field('ยี่ห้อ'),field('รุ่น'),field('ชนิดอ้อย (อ้อยลำ/อ้อยตัด)')] + [field(x,'check') for x in ['ล้อรถ','โบลต์ล็อคดั๊มพ์','ฝาท้าย','การ์ดกันหม้อน้ำ','โซ่/สลิงผ้า']],
 'FM-ML02-ML-01': [field('หัวข้องาน'),field('สัปดาห์ที่'),field('วันที่เริ่ม','date'),field('กำหนดเสร็จ','date'),field('จำนวนแรงงาน','number','คน'),field('ชั่วโมงปกติ','number'),field('ชั่วโมง OT','number'),field('ความก้าวหน้า/ผลการซ่อม')],
 'FM-ML02-ML-03': [field('ชื่อสาร/ยี่ห้อ'),field('Lot No./Batch No.'),field('เริ่มใช้/รับเข้า','datetime-local'),field('หยุดใช้','datetime-local'),field('ส่งเข้าอาคารขยะ','datetime-local'),field('หน่วย')] + numbers('รับเข้า|ใช้จริง|คงเหลือ|ภาชนะรับเข้า|ภาชนะส่งเข้าอาคารขยะ|ภาชนะคงเหลือ'),
 'FM-PD01-PD-01': [field(x,'time') for x in ['เริ่มหยุดหีบ','เริ่มลดโฟล/ปริมาณ','หม้อต้มพร้อม','ลูกหีบพร้อม','เรียกรถอ้อยเข้าดั๊มพ์','เริ่มหีบ','เริ่มเพิ่มโฟล/ปริมาณ']],
 'FM-PD01-PD-02': [field('กะ'),field('สถานที่/แผนก'),field('ปัญหา/ชำรุด'),field('สาเหตุ'),field('วิธีการแก้ไข'),field('ช่วงเวลา'),field('ลด/หยุด'),field('งานค้าง/ส่งต่อ')],
 'FM-PD02-PD-02': [field('ชื่อวัสดุ/รหัสตามทะเบียน'),field('จุดติดตั้ง'),field('จำนวน','number','หน่วย'),field('ความถี่'),field('สภาพ/ผลตรวจ','check'),field('รายละเอียดแตกชำรุด/แก้ไข')],
}
schemas['FM-ML01-ML-08'] += schemas['FM-ML01-ML-07']

items=[]
for f in sorted(SOURCE.rglob('*')):
 if not f.is_file(): continue
 data=f.read_bytes(); code=re.search(r'FM-[A-Z]+\d+-[A-Z]+-\d+',f.name)
 code=code.group() if code else f.stem
 rev=re.search(r'(?:Rev|Rve)\.?\s*(\d+)',f.name,re.I)
 pages=[]
 if f.suffix.lower()=='.pdf':
  reader=PdfReader(f); pages=[p.extract_text() or '' for p in reader.pages]
 else:
  wb=load_workbook(f,read_only=True,data_only=True)
  pages=['ชีต: '+s.title+'\n'+'\n'.join(' | '.join(str(v) if v is not None else '' for v in row) for row in s.iter_rows(values_only=True) if any(v is not None for v in row)) for s in wb]
  wb.close()
 fields=schemas.get(code,[])
 if str(f.relative_to(SOURCE)).startswith('01.') and f.suffix.lower()=='.pdf':
  layout=reader.pages[0].extract_text(extraction_mode='layout')
  expected=1; fields=[]
  for line in layout.splitlines():
   m=re.search(r'(?:^|\s)(\d{1,2})\s{2,}([^\d\s].*?)\s{2,}(.+?)\s*$',line)
   if m and int(m[1])==expected:
    fields.append(field(m[2].strip(),'check',standard=m[3].strip()));expected+=1
  if not fields: raise ValueError('No tool checklist extracted: '+f.name)
  if code=='FM-SF04-SF-14':
   labels='รางเครน/รางวิ่ง|สะพานเครน|บีมหัวท้ายและชุดล้อ|มอเตอร์ขับเคลื่อนเครนตามแนวราง|ฮอยท์และทรอลเลย์ (ชุดกว้าน)|ชุดรอกและตะขอยก|ขาล็อคตะขอยก 6.1|ระบบควบคุมเครน|เบรกกว้าน 7.1|สวิทซ์ควบคุมช่วงยก/โหลดเกิน 7.2|แผ่นประกับ cheek weights|กลไกบังคับร่องสลิง|พิกัดยก|สัญญาณเตือน'
   fields=[field(x,'check',standard='ตรวจตามรายละเอียดรายการนี้ในต้นฉบับ') for x in labels.split('|')]
  if code=='FM-SF04-SF-21':
   fields=[field(x,'check',standard=s) for x,s in [('สายไฟฟ้า','ไม่ชำรุด ฉนวนสายไฟสภาพปกติ'),('การ์ดป้องกันความร้อน','ไม่แตกชำรุด น็อตยึดไม่คลายตัว'),('มือจับด้านข้าง','ไม่แตกชำรุด ด้ามไม่คลอน'),('สวิทช์เปิด/ปิด','ไม่ชำรุด สวิตช์ปิดเปิดได้'),('สภาพโดยรวม','ตัวเครื่องไม่บิดเบี้ยว ไม่คลอน')]]
  if code=='FM-SF04-SF-23': fields.append(field('อุปกรณ์ดูดซับแรงกระชาก','check',standard='ไม่หลุดลุ่ย ฉีกขาด'))
  for x in fields: x['standard'] += ' — ตรวจรายละเอียดต่อเนื่องและข้อควรระวังในต้นฉบับ'
 title=re.sub(r'^.*?'+re.escape(code)+r'\s*(?:Rev|Rve)\.?\s*\d+\s*','',f.stem,flags=re.I) if code.startswith('FM-') else f.stem
 items.append(dict(id=f'doc-{len(items)+1:03}',code=code,revision=rev[1] if rev else '',title=title,path=str(f.relative_to(SOURCE)),extension=f.suffix.lower(),sha256=hashlib.sha256(data).hexdigest(),pages=pages,fields=fields,mode='form' if fields else 'reference',group='เครื่องมือ' if str(f.relative_to(SOURCE)).startswith('01.') else 'ใบอนุญาตทำงาน' if str(f.relative_to(SOURCE)).startswith('02.') else 'ลูกหีบและเอกสารอื่น'))
out=ROOT/'public/assets/source-forms';out.mkdir(parents=True,exist_ok=True)
(out/'catalogue.js').write_text('globalThis.SourceDocumentsRoot = '+json.dumps(str(SOURCE),ensure_ascii=False)+';\nglobalThis.SourceCatalogue = '+json.dumps(items,ensure_ascii=False,indent=2)+';\n',encoding='utf-8')
(ROOT/'logs/source-form-review/manifest.json').write_text(json.dumps([dict(id=e['id'],path=e['path'],sha256=e['sha256'],fields=len(e['fields'])) for e in items],ensure_ascii=False,indent=2),encoding='utf-8')
print('Documents:',len(items),'Forms:',sum(bool(e['fields']) for e in items),'Tool checklist counts:',[len(e['fields']) for e in items if e['group']=='เครื่องมือ'])
